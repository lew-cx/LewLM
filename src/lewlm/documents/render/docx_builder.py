"""DOCX document renderer."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from lewlm.core.errors import DocumentGenerationError
from lewlm.documents.ir.models import (
    CalloutBlock,
    DocumentIR,
    DocumentOutputFormat,
    ImageBlock,
    ListBlock,
    ParagraphBlock,
    TableBlock,
)
from lewlm.documents.ir.style import (
    DocumentStyleSheet,
    ResolvedStyle,
    StyleElement,
    element_for_block,
    hex_digits,
)
from lewlm.documents.render.base import DocumentRenderer
from lewlm.documents.validators.ir import DocumentIRValidator


class DocxDocumentRenderer(DocumentRenderer):
    """Render DOCX.

    Colour, font family, point size, and emphasis tokens are applied to runs;
    surface and accent colours become paragraph and table-cell shading.
    """

    output_format = DocumentOutputFormat.DOCX
    media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    file_extension = ".docx"

    def __init__(self, validator: DocumentIRValidator) -> None:
        self.validator = validator

    def render(self, document: DocumentIR) -> bytes:
        self.validator.validate(document)
        try:
            from docx import Document as WordDocument
        except ImportError as exc:
            raise DocumentGenerationError(
                "DOCX generation requires the `python-docx` dependency.",
                details={"renderer": self.output_format.value},
            ) from exc

        stylesheet = DocumentStyleSheet.from_document(document)
        word_document = WordDocument()
        word_document.core_properties.title = document.title
        self._render_header_footer(word_document, document, stylesheet)
        title_style = stylesheet.for_element(StyleElement.HEADING)
        _style_paragraph(word_document.add_heading(document.title, level=0), title_style)

        for section in document.sections:
            heading_style = stylesheet.for_element(StyleElement.HEADING, token_names=section.style_tokens)
            if section.heading:
                _style_paragraph(
                    word_document.add_heading(section.heading, level=max(1, min(section.level, 9))),
                    heading_style,
                )
            for block in section.blocks:
                resolved = stylesheet.for_element(
                    element_for_block(block),
                    token_names=[*section.style_tokens, *block.style_tokens],
                )
                if isinstance(block, ParagraphBlock):
                    _style_paragraph(word_document.add_paragraph(block.text), resolved)
                elif isinstance(block, ListBlock):
                    style = "List Number" if block.ordered else "List Bullet"
                    for item in block.items:
                        _style_paragraph(word_document.add_paragraph(item, style=style), resolved)
                elif isinstance(block, TableBlock):
                    self._render_table(
                        word_document,
                        block,
                        resolved,
                        stylesheet.accent_color_for(
                            StyleElement.TABLE,
                            token_names=[*section.style_tokens, *block.style_tokens],
                        ),
                    )
                elif isinstance(block, CalloutBlock):
                    paragraph = word_document.add_paragraph()
                    prefix = f"{block.kind.upper()}"
                    if block.title:
                        prefix = f"{prefix}: {block.title}"
                    paragraph.add_run(prefix).bold = True
                    paragraph.add_run(f" {block.body}")
                    _style_paragraph(paragraph, resolved)
                elif isinstance(block, ImageBlock):
                    self._render_image(word_document, block, resolved)

        if document.citations:
            word_document.add_heading(document.references_title, level=1)
            for citation in document.citations:
                resolved = stylesheet.for_element(StyleElement.CITATION, token_names=citation.style_tokens)
                text = citation.text if citation.url is None else f"{citation.text} ({citation.url})"
                _style_paragraph(word_document.add_paragraph(f"[{citation.label}] {text}"), resolved)

        buffer = BytesIO()
        word_document.save(buffer)
        return buffer.getvalue()

    def _render_header_footer(self, word_document, document: DocumentIR, stylesheet: DocumentStyleSheet) -> None:
        section = word_document.sections[0]
        if document.header:
            header = section.header.paragraphs[0]
            header.text = " | ".join(value for value in (document.header.left, document.header.center, document.header.right) if value)
            _style_paragraph(header, stylesheet.for_element(StyleElement.HEADER, token_names=document.header.style_tokens))
        if document.footer:
            footer = section.footer.paragraphs[0]
            footer.text = " | ".join(value for value in (document.footer.left, document.footer.center, document.footer.right) if value)
            _style_paragraph(footer, stylesheet.for_element(StyleElement.FOOTER, token_names=document.footer.style_tokens))

    def _render_table(self, word_document, block: TableBlock, style: ResolvedStyle, accent_color: str | None) -> None:
        column_count = len(block.headers) if block.headers else len(block.rows[0])
        table = word_document.add_table(rows=0, cols=column_count)
        table.style = "Table Grid"
        if block.headers:
            header_cells = table.add_row().cells
            for index, header in enumerate(block.headers):
                header_cells[index].text = header
                if accent_color is not None:
                    _shade_element(header_cells[index]._tc.get_or_add_tcPr(), accent_color)
                for paragraph in header_cells[index].paragraphs:
                    _style_paragraph(paragraph, style, shade=False)
        for row in block.rows:
            row_cells = table.add_row().cells
            for index, value in enumerate(row):
                row_cells[index].text = value
                for paragraph in row_cells[index].paragraphs:
                    _style_paragraph(paragraph, style, shade=False)
        if block.caption:
            _style_paragraph(word_document.add_paragraph(block.caption), style)

    def _render_image(self, word_document, block: ImageBlock, style: ResolvedStyle) -> None:
        if block.path is not None:
            word_document.add_picture(str(Path(block.path).expanduser()))
        else:
            _style_paragraph(word_document.add_paragraph(f"[Image] {block.alt_text}"), style)
        if block.caption:
            _style_paragraph(word_document.add_paragraph(block.caption), style)


def _style_paragraph(paragraph, style: ResolvedStyle, *, shade: bool = True) -> None:
    """Apply a resolved style to every run of a paragraph."""

    if style.is_empty:
        return
    from docx.shared import Pt, RGBColor

    for run in paragraph.runs:
        if style.text_color is not None:
            run.font.color.rgb = RGBColor.from_string(hex_digits(style.text_color))
        if style.font_family is not None:
            run.font.name = style.font_family
        if style.font_size_pt is not None:
            run.font.size = Pt(style.font_size_pt)
        if style.bold:
            run.bold = True
        if style.italic:
            run.italic = True
    if shade and style.background_color is not None:
        _shade_element(paragraph._p.get_or_add_pPr(), style.background_color)


def _shade_element(properties, color: str) -> None:
    """Attach a solid `w:shd` fill to a paragraph or table-cell property element."""

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_digits(color))
    properties.append(shading)
