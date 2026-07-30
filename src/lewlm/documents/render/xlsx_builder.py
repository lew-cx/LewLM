"""XLSX document renderer."""

from __future__ import annotations

from io import BytesIO

from lewlm.core.errors import DocumentGenerationError
from lewlm.documents.ir.models import (
    CalloutBlock,
    DocumentIR,
    DocumentOutputFormat,
    ImageBlock,
    ListBlock,
    ParagraphBlock,
    TableBlock,
)
from lewlm.documents.ir.style import (
    DocumentStyleSheet,
    ResolvedStyle,
    StyleElement,
    element_for_block,
    hex_digits,
)
from lewlm.documents.render.base import DocumentRenderer
from lewlm.documents.validators.ir import DocumentIRValidator


class XlsxDocumentRenderer(DocumentRenderer):
    """Render XLSX.

    Colour, font family, point size, and emphasis tokens become cell fonts;
    surface and accent colours become solid cell fills.
    """

    output_format = DocumentOutputFormat.XLSX
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    file_extension = ".xlsx"

    def __init__(self, validator: DocumentIRValidator) -> None:
        self.validator = validator

    def render(self, document: DocumentIR) -> bytes:
        self.validator.validate(document)
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError as exc:
            raise DocumentGenerationError(
                "XLSX generation requires the `openpyxl` dependency.",
                details={"renderer": self.output_format.value},
            ) from exc

        stylesheet = DocumentStyleSheet.from_document(document)
        title_style = stylesheet.for_element(StyleElement.HEADING)
        workbook = Workbook()
        first_sheet = workbook.active
        for index, section in enumerate(document.sections):
            heading_style = stylesheet.for_element(StyleElement.HEADING, token_names=section.style_tokens)
            worksheet = first_sheet if index == 0 else workbook.create_sheet()
            worksheet.title = self._sheet_title(section.heading or f"Section {index + 1}")
            row = 1
            if index == 0:
                cell = worksheet.cell(row=row, column=1, value=document.title)
                cell.font = Font(bold=True)
                _style_cell(cell, title_style)
                row += 2
            if section.heading:
                cell = worksheet.cell(row=row, column=1, value=section.heading)
                cell.font = Font(bold=True)
                _style_cell(cell, heading_style)
                row += 2
            for block in section.blocks:
                resolved = stylesheet.for_element(
                    element_for_block(block),
                    token_names=[*section.style_tokens, *block.style_tokens],
                )
                row = self._render_block(
                    worksheet,
                    block,
                    row,
                    resolved,
                    stylesheet.accent_color_for(
                        element_for_block(block),
                        token_names=[*section.style_tokens, *block.style_tokens],
                    ),
                )
                row += 1

        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def _render_block(self, worksheet, block, row: int, style: ResolvedStyle, accent_color: str | None) -> int:
        if isinstance(block, ParagraphBlock):
            _style_cell(worksheet.cell(row=row, column=1, value=block.text), style)
            return row
        if isinstance(block, ListBlock):
            for index, item in enumerate(block.items, start=1):
                prefix = f"{index}. " if block.ordered else "- "
                _style_cell(worksheet.cell(row=row, column=1, value=f"{prefix}{item}"), style)
                row += 1
            return row - 1
        if isinstance(block, CalloutBlock):
            prefix = block.kind.upper()
            if block.title:
                prefix = f"{prefix}: {block.title}"
            _style_cell(worksheet.cell(row=row, column=1, value=f"{prefix} {block.body}"), style)
            return row
        if isinstance(block, ImageBlock):
            _style_cell(worksheet.cell(row=row, column=1, value=f"[Image] {block.alt_text}"), style)
            if block.caption:
                _style_cell(worksheet.cell(row=row + 1, column=1, value=block.caption), style)
                return row + 1
            return row
        if isinstance(block, TableBlock):
            headers = block.headers or [f"column_{index + 1}" for index in range(len(block.rows[0]))]
            if block.caption:
                _style_cell(worksheet.cell(row=row, column=1, value=block.caption), style)
                row += 1
            for column, header in enumerate(headers, start=1):
                cell = worksheet.cell(row=row, column=column, value=header)
                _style_cell(cell, style)
                if accent_color is not None:
                    _fill_cell(cell, accent_color)
            row += 1
            for values in block.rows:
                for column, value in enumerate(values, start=1):
                    _style_cell(worksheet.cell(row=row, column=column, value=value), style)
                row += 1
            return row - 1
        return row

    def _sheet_title(self, value: str) -> str:
        sanitized = "".join(char for char in value if char not in {"\\", "/", "*", "[", "]", ":", "?"}).strip()
        return (sanitized or "Sheet")[:31]


def _style_cell(cell, style: ResolvedStyle) -> None:
    """Apply a resolved style to one worksheet cell, preserving existing font traits."""

    if style.is_empty:
        return
    from openpyxl.styles import Font

    current = cell.font
    cell.font = Font(
        name=style.font_family or current.name,
        size=style.font_size_pt or current.size,
        bold=style.bold or bool(current.bold),
        italic=style.italic or bool(current.italic),
        color=f"FF{hex_digits(style.text_color)}" if style.text_color is not None else current.color,
    )
    if style.background_color is not None:
        _fill_cell(cell, style.background_color)


def _fill_cell(cell, color: str) -> None:
    from openpyxl.styles import PatternFill

    digits = f"FF{hex_digits(color)}"
    cell.fill = PatternFill(fill_type="solid", start_color=digits, end_color=digits)
