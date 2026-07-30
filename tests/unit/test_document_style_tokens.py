from __future__ import annotations

from io import BytesIO

import pytest

pytest.importorskip("openpyxl")
pytest.importorskip("docx")

import docx
from openpyxl import load_workbook

from lewlm.core.errors import DocumentValidationError
from lewlm.documents.ir.models import (
    CalloutBlock,
    Citation,
    DocumentIR,
    DocumentOutputFormat,
    DocumentSection,
    HeaderFooterContent,
    ListBlock,
    ParagraphBlock,
    StyleToken,
    TableBlock,
)
from lewlm.documents.ir.style import (
    DocumentStyleSheet,
    StyleElement,
    StyleTokenRole,
    normalize_style_value,
    style_role_specs,
)
from lewlm.documents.render.pdf_builder import PdfDocumentRenderer
from lewlm.documents.service import DocumentGenerationService
from lewlm.documents.validators.ir import DocumentIRValidator

STYLED_FORMATS = (
    DocumentOutputFormat.JSON,
    DocumentOutputFormat.DOCX,
    DocumentOutputFormat.PDF,
    DocumentOutputFormat.XLSX,
)
INERT_FORMATS = (
    DocumentOutputFormat.TEXT,
    DocumentOutputFormat.MARKDOWN,
    DocumentOutputFormat.CSV,
)


def _branded_document(accent: str = "#FF0000", extra_tokens: list[StyleToken] | None = None) -> DocumentIR:
    return DocumentIR(
        title="Branded Report",
        style_tokens=[
            StyleToken(name="accent_color", value=accent, applies_to="document"),
            StyleToken(name="heading_color", value="#111827"),
            StyleToken(name="body_color", value="#374151"),
            StyleToken(name="body_font", value="Georgia"),
            StyleToken(name="body_font_size_pt", value="11"),
            StyleToken(name="emphasis", value="bold_italic"),
            *(extra_tokens or []),
        ],
        header=HeaderFooterContent(left="Acme"),
        footer=HeaderFooterContent(right="Internal"),
        sections=[
            DocumentSection(
                heading="Summary",
                blocks=[
                    ParagraphBlock(text="Plain body."),
                    ParagraphBlock(text="Emphasised body.", style_tokens=["emphasis"]),
                    ListBlock(items=["one", "two"]),
                    CalloutBlock(kind="info", title="Note", body="Callout body."),
                ],
            ),
            DocumentSection(
                heading="Budget",
                blocks=[TableBlock(headers=["Category", "Amount"], rows=[["Hosting", "1200"]], caption="Spend")],
            ),
        ],
        citations=[Citation(label="1", text="Finance worksheet")],
    )


@pytest.mark.parametrize("output_format", STYLED_FORMATS)
def test_accent_token_value_changes_styled_artifacts(output_format: DocumentOutputFormat) -> None:
    service = DocumentGenerationService()

    red = service.generate(_branded_document("#FF0000"), output_format=output_format).content
    blue = service.generate(_branded_document("#0000FF"), output_format=output_format).content

    assert red != blue


@pytest.mark.parametrize("output_format", INERT_FORMATS)
def test_colour_tokens_are_inert_in_unstyled_formats(output_format: DocumentOutputFormat) -> None:
    service = DocumentGenerationService()

    red = service.generate(_branded_document("#FF0000"), output_format=output_format).content
    blue = service.generate(_branded_document("#0000FF"), output_format=output_format).content

    assert red == blue


def test_markdown_renders_emphasis_tokens_only() -> None:
    service = DocumentGenerationService()

    markdown = service.generate(_branded_document(), output_format=DocumentOutputFormat.MARKDOWN).content.decode("utf-8")

    assert "***Emphasised body.***" in markdown
    assert "\nPlain body." in markdown
    assert "#FF0000" not in markdown


def test_docx_applies_colour_font_size_and_accent_shading() -> None:
    service = DocumentGenerationService()

    artifact = service.generate(_branded_document(), output_format=DocumentOutputFormat.DOCX)
    word_document = docx.Document(BytesIO(artifact.content))

    paragraphs = {paragraph.text: paragraph for paragraph in word_document.paragraphs}
    heading_run = paragraphs["Summary"].runs[0]
    assert str(heading_run.font.color.rgb) == "111827"
    body_run = paragraphs["Plain body."].runs[0]
    assert str(body_run.font.color.rgb) == "374151"
    assert body_run.font.name == "Georgia"
    assert body_run.font.size.pt == 11
    emphasised_run = paragraphs["Emphasised body."].runs[0]
    assert emphasised_run.bold is True
    assert emphasised_run.italic is True

    header_cell_xml = word_document.tables[0].rows[0].cells[0]._tc.xml
    assert 'w:fill="FF0000"' in header_cell_xml


def test_xlsx_applies_font_and_accent_fill() -> None:
    service = DocumentGenerationService()

    artifact = service.generate(_branded_document(), output_format=DocumentOutputFormat.XLSX)
    workbook = load_workbook(BytesIO(artifact.content))

    header_cell = workbook["Budget"]["A4"]
    assert header_cell.value == "Category"
    assert header_cell.fill.start_color.rgb == "FFFF0000"
    assert header_cell.font.name == "Georgia"
    assert header_cell.font.size == 11
    assert workbook["Budget"]["A5"].font.color.rgb == "FF374151"


def test_pdf_html_path_emits_validated_declarations_only() -> None:
    document = _branded_document()
    renderer = PdfDocumentRenderer(DocumentIRValidator())

    markup = renderer._to_html(document, DocumentStyleSheet.from_document(document))

    assert "border-left: 4px solid #FF0000" in markup
    assert '<th style="background-color: #FF0000">' in markup
    assert "color: #111827" in markup
    assert "font-family: 'Georgia'; font-size: 11pt" in markup
    assert "font-weight: bold; font-style: italic" in markup


def test_lineage_tokens_are_preserved_but_never_rendered() -> None:
    service = DocumentGenerationService()
    lineage = StyleToken(name="ocr", value="tesseract-provenance", applies_to="pipeline")

    without = service.generate(_branded_document(), output_format=DocumentOutputFormat.DOCX).content
    with_lineage = service.generate(
        _branded_document(extra_tokens=[lineage]),
        output_format=DocumentOutputFormat.DOCX,
    ).content
    json_bytes = service.generate(
        _branded_document(extra_tokens=[lineage]),
        output_format=DocumentOutputFormat.JSON,
    ).content

    assert len(without) == len(with_lineage)
    assert b"tesseract-provenance" in json_bytes


def test_block_references_to_undeclared_tokens_are_ignored() -> None:
    service = DocumentGenerationService()
    document = DocumentIR(
        title="Ingested",
        sections=[
            DocumentSection(
                heading="Body",
                blocks=[ParagraphBlock(text="print(1)", style_tokens=["code", "Quote", "emphasis"])],
            ),
        ],
    )

    markdown = service.generate(document, output_format=DocumentOutputFormat.MARKDOWN).content.decode("utf-8")

    assert "print(1)" in markdown
    assert "**" not in markdown


def test_section_token_references_cascade_to_blocks() -> None:
    service = DocumentGenerationService()
    document = DocumentIR(
        title="Cascade",
        style_tokens=[StyleToken(name="emphasis", value="bold")],
        sections=[
            DocumentSection(
                heading="Body",
                style_tokens=["emphasis"],
                blocks=[ParagraphBlock(text="inherited")],
            ),
        ],
    )

    markdown = service.generate(document, output_format=DocumentOutputFormat.MARKDOWN).content.decode("utf-8")

    assert "## **Body**" in markdown
    assert "**inherited**" in markdown


def test_applies_to_restricts_a_token_to_one_element_class() -> None:
    document = DocumentIR(
        title="Scoped",
        style_tokens=[StyleToken(name="body_color", value="#AA0000", applies_to="callout")],
        sections=[DocumentSection(blocks=[ParagraphBlock(text="body")])],
    )
    stylesheet = DocumentStyleSheet.from_document(document)

    assert stylesheet.for_element(StyleElement.CALLOUT).text_color == "#AA0000"
    assert stylesheet.for_element(StyleElement.PARAGRAPH).text_color is None


def test_accent_scope_does_not_leak_into_table_or_callout_defaults() -> None:
    document = DocumentIR(
        title="Scoped accent",
        style_tokens=[StyleToken(name="accent_color", value="#123456", applies_to="paragraph")],
        sections=[DocumentSection(blocks=[ParagraphBlock(text="body")])],
    )
    stylesheet = DocumentStyleSheet.from_document(document)

    assert stylesheet.accent_color_for(StyleElement.PARAGRAPH) == "#123456"
    assert stylesheet.accent_color_for(StyleElement.TABLE) is None
    assert stylesheet.accent_color_for(StyleElement.CALLOUT) is None


def test_colour_values_are_normalized_to_upper_case() -> None:
    document = DocumentIR(
        title="Normalized",
        style_tokens=[StyleToken(name="heading_color", value="#aabbcc")],
        sections=[DocumentSection(blocks=[ParagraphBlock(text="body")])],
    )

    stylesheet = DocumentStyleSheet.from_document(document)

    assert stylesheet.for_element(StyleElement.HEADING).text_color == "#AABBCC"


@pytest.mark.parametrize(
    ("name", "value"),
    [
        ("accent_color", "red"),
        ("accent_color", "#FFF"),
        ("heading_color", "#GGGGGG"),
        ("body_font", "Georgia; color: red"),
        ("body_font", "'Comic Sans'"),
        ("body_font_size_pt", "huge"),
        ("body_font_size_pt", "500"),
        ("body_font_size_pt", "0"),
        ("emphasis", "underline"),
    ],
)
def test_reserved_tokens_reject_invalid_values(name: str, value: str) -> None:
    service = DocumentGenerationService()
    document = DocumentIR(
        title="Invalid",
        style_tokens=[StyleToken(name=name, value=value)],
        sections=[DocumentSection(blocks=[ParagraphBlock(text="body")])],
    )

    with pytest.raises(DocumentValidationError) as error:
        service.generate(document, output_format=DocumentOutputFormat.MARKDOWN)

    assert error.value.details["name"] == name


def test_reserved_tokens_reject_unknown_element_scopes() -> None:
    service = DocumentGenerationService()
    document = DocumentIR(
        title="Invalid scope",
        style_tokens=[StyleToken(name="accent_color", value="#FF0000", applies_to="sidebar")],
        sections=[DocumentSection(blocks=[ParagraphBlock(text="body")])],
    )

    with pytest.raises(DocumentValidationError) as error:
        service.generate(document, output_format=DocumentOutputFormat.MARKDOWN)

    assert error.value.details["applies_to"] == "sidebar"


def test_duplicate_style_token_names_are_rejected() -> None:
    service = DocumentGenerationService()
    document = DocumentIR(
        title="Duplicate",
        style_tokens=[
            StyleToken(name="accent_color", value="#FF0000"),
            StyleToken(name="accent_color", value="#0000FF"),
        ],
        sections=[DocumentSection(blocks=[ParagraphBlock(text="body")])],
    )

    with pytest.raises(DocumentValidationError) as error:
        service.generate(document, output_format=DocumentOutputFormat.MARKDOWN)

    assert error.value.details["name"] == "accent_color"


def test_every_reserved_role_has_a_published_spec() -> None:
    specs = style_role_specs()

    assert {spec.role for spec in specs} == set(StyleTokenRole)
    for spec in specs:
        assert spec.description


def test_normalize_style_value_trims_and_normalizes() -> None:
    assert normalize_style_value(StyleTokenRole.ACCENT_COLOR, "  #ff0000 ") == "#FF0000"
    assert normalize_style_value(StyleTokenRole.BODY_FONT_SIZE_PT, "11.0") == "11"
    assert normalize_style_value(StyleTokenRole.EMPHASIS, "bold") == "bold"
